"""Conversie- en validatielogica voor peilfilterloggerbestanden."""

from __future__ import annotations

import csv
import io
import re
import struct
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from statistics import mean
from typing import Iterable, Iterator

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

CMH2O_TO_HPA = 0.980665
DEFAULT_PRESSURE_SCALE = 0.1
DEFAULT_PRESSURE_OFFSET_HPA = 0.0
EXCEL_MAX_ROWS = 1_048_576
RECORD_STRUCT = struct.Struct("<HH")
DATETIME_PATTERN = re.compile(
    rb"(?P<time>\d{2}:\d{2}:\d{2})\s+(?P<date>\d{2}/\d{2}/\d{2})"
)
INTERVAL_PATTERN = re.compile(
    rb"(?P<days>\d{2})\s+(?P<hours>\d{2}):(?P<minutes>\d{2}):"
    rb"(?P<seconds>\d{2})\s+\dT"
)


@dataclass(frozen=True)
class LoggerMetadata:
    """Metadata uit de loggerheader."""

    start_datetime: datetime
    sample_interval: timedelta
    data_offset: int
    header_text: str
    detected_datetimes: tuple[datetime, ...]


@dataclass(frozen=True)
class Measurement:
    """Eén geconverteerde meting."""

    timestamp: datetime
    pressure_hpa: float
    raw_pressure: int
    raw_temperature: int


@dataclass(frozen=True)
class ReferenceMeasurement:
    """Eén referentiemeting uit een geëxporteerd tekstbestand."""

    timestamp: datetime
    pressure_hpa: float


@dataclass(frozen=True)
class DifferenceRow:
    """Verschillen tussen één berekende en één referentiemeting."""

    index: int
    calculated_timestamp: datetime
    reference_timestamp: datetime
    time_difference_seconds: float
    calculated_pressure_hpa: float
    reference_pressure_hpa: float
    pressure_difference_hpa: float


@dataclass(frozen=True)
class ValidationSummary:
    """Samenvatting van een vergelijking met referentiedata."""

    compared_count: int
    calculated_count: int
    reference_count: int
    mean_pressure_difference_hpa: float
    mean_absolute_pressure_difference_hpa: float
    maximum_absolute_pressure_difference_hpa: float
    maximum_absolute_time_difference_seconds: float
    pressure_within_tolerance: bool
    time_within_tolerance: bool
    rows: tuple[DifferenceRow, ...]


@dataclass(frozen=True)
class ExcelResult:
    """Gegenereerd Excel-bestand plus conversieresultaten."""

    filename: str
    content: bytes
    measurements: tuple[Measurement, ...]
    metadata: LoggerMetadata
    validation: ValidationSummary | None


def excel_serial_to_datetime(value: float) -> datetime:
    """Converteer een Excel-seriële datum volgens het 1900-datumsysteem."""
    return datetime(1899, 12, 30) + timedelta(days=value)


def parse_logger_datetime(time_bytes: bytes, date_bytes: bytes) -> datetime:
    """Converteer DD/MM/YY en HH:MM:SS naar een datetime."""
    value = f"{date_bytes.decode('ascii')} {time_bytes.decode('ascii')}"
    return datetime.strptime(value, "%d/%m/%y %H:%M:%S")


def parse_interval(header: bytes) -> timedelta:
    """Lees en valideer de meetinterval uit de header."""
    match = INTERVAL_PATTERN.search(header)
    if match is None:
        raise ValueError("De meetinterval kon niet uit de header worden gelezen.")
    interval = timedelta(
        days=int(match.group("days")),
        hours=int(match.group("hours")),
        minutes=int(match.group("minutes")),
        seconds=int(match.group("seconds")),
    )
    if interval <= timedelta(0):
        raise ValueError("De meetinterval moet groter zijn dan nul.")
    return interval


def detect_metadata(
    file_data: bytes,
    absolute_data_offset: int | None = None,
    start_datetime_override: datetime | None = None,
    datetime_field_index: int = 0,
) -> LoggerMetadata:
    """Detecteer metadata en pas expliciete gebruikerscorrecties toe."""
    if not file_data:
        raise ValueError("Het invoerbestand is leeg.")
    matches = list(DATETIME_PATTERN.finditer(file_data[:1024]))
    if not matches:
        raise ValueError("Geen datum/tijdvelden gevonden in de eerste 1024 bytes.")
    detected = tuple(
        parse_logger_datetime(match.group("time"), match.group("date"))
        for match in matches
    )
    if datetime_field_index < 0 or datetime_field_index >= len(detected):
        raise ValueError(
            f"Datumveld {datetime_field_index + 1} bestaat niet; "
            f"gevonden: {len(detected)}."
        )
    start_datetime = start_datetime_override or detected[datetime_field_index]
    if absolute_data_offset is not None:
        data_offset = absolute_data_offset
    elif len(matches) >= 2:
        data_offset = matches[1].end()
    else:
        raise ValueError(
            "Het einde van de header kon niet automatisch worden bepaald. "
            "Vul een handmatige byte-offset in."
        )
    if data_offset < 0 or data_offset >= len(file_data):
        raise ValueError(
            f"Ongeldige data-offset {data_offset}; "
            f"bestandsgrootte is {len(file_data)} bytes."
        )
    header = file_data[:data_offset]
    return LoggerMetadata(
        start_datetime=start_datetime,
        sample_interval=parse_interval(header),
        data_offset=data_offset,
        header_text=header.decode("cp1252", errors="replace"),
        detected_datetimes=detected,
    )


def raw_pressure_to_hpa(
    raw_pressure: int,
    pressure_scale: float,
    pressure_offset_hpa: float = DEFAULT_PRESSURE_OFFSET_HPA,
) -> float:
    """Converteer ruwe druk via cmH2O naar hPa en pas een offset toe."""
    if pressure_scale <= 0:
        raise ValueError("De drukschaalfactor moet groter zijn dan nul.")
    return raw_pressure * pressure_scale * CMH2O_TO_HPA + pressure_offset_hpa


def iter_measurements(
    file_data: bytes,
    metadata: LoggerMetadata,
    pressure_scale: float = DEFAULT_PRESSURE_SCALE,
    pressure_offset_hpa: float = DEFAULT_PRESSURE_OFFSET_HPA,
) -> Iterator[Measurement]:
    """Itereer over alle volledige binaire meetrecords."""
    binary_data = file_data[metadata.data_offset:]
    for index in range(len(binary_data) // RECORD_STRUCT.size):
        raw_pressure, raw_temperature = RECORD_STRUCT.unpack_from(
            binary_data, index * RECORD_STRUCT.size
        )
        yield Measurement(
            timestamp=metadata.start_datetime + index * metadata.sample_interval,
            pressure_hpa=raw_pressure_to_hpa(
                raw_pressure, pressure_scale, pressure_offset_hpa
            ),
            raw_pressure=raw_pressure,
            raw_temperature=raw_temperature,
        )


def parse_reference_data(text: str) -> tuple[ReferenceMeasurement, ...]:
    """Lees tab-, puntkomma- of spatiegescheiden Excel-datum- en hPa-kolommen."""
    measurements: list[ReferenceMeasurement] = []
    for line_number, raw_line in enumerate(text.splitlines(), start=1):
        line = raw_line.strip()
        if not line:
            continue
        parts = [part for part in re.split(r"\t+|;|\s+", line) if part]
        if len(parts) < 2:
            raise ValueError(f"Referentieregel {line_number} bevat minder dan twee kolommen.")
        try:
            serial = float(parts[0].replace(",", "."))
            pressure = float(parts[1].replace(",", "."))
        except ValueError as exc:
            if line_number == 1:
                continue
            raise ValueError(
                f"Referentieregel {line_number} bevat geen geldige getallen."
            ) from exc
        measurements.append(ReferenceMeasurement(excel_serial_to_datetime(serial), pressure))
    if not measurements:
        raise ValueError("De referentiedata bevat geen geldige meetregels.")
    return tuple(measurements)


def compare_measurements(
    calculated: Iterable[Measurement],
    reference: Iterable[ReferenceMeasurement],
    pressure_tolerance_hpa: float = 0.1,
    time_tolerance_seconds: float = 1.0,
) -> ValidationSummary:
    """Vergelijk berekende en geëxporteerde meetwaarden op dezelfde rij-index."""
    if pressure_tolerance_hpa < 0 or time_tolerance_seconds < 0:
        raise ValueError("Toleranties mogen niet negatief zijn.")
    calculated_values = tuple(calculated)
    reference_values = tuple(reference)
    rows: list[DifferenceRow] = []
    for index, (actual, expected) in enumerate(
        zip(calculated_values, reference_values), start=1
    ):
        rows.append(
            DifferenceRow(
                index=index,
                calculated_timestamp=actual.timestamp,
                reference_timestamp=expected.timestamp,
                time_difference_seconds=(
                    actual.timestamp - expected.timestamp
                ).total_seconds(),
                calculated_pressure_hpa=actual.pressure_hpa,
                reference_pressure_hpa=expected.pressure_hpa,
                pressure_difference_hpa=actual.pressure_hpa - expected.pressure_hpa,
            )
        )
    if not rows:
        raise ValueError("Er zijn geen meetregels om met de referentie te vergelijken.")
    pressure_differences = [row.pressure_difference_hpa for row in rows]
    time_differences = [row.time_difference_seconds for row in rows]
    max_pressure = max(abs(value) for value in pressure_differences)
    max_time = max(abs(value) for value in time_differences)
    return ValidationSummary(
        compared_count=len(rows),
        calculated_count=len(calculated_values),
        reference_count=len(reference_values),
        mean_pressure_difference_hpa=mean(pressure_differences),
        mean_absolute_pressure_difference_hpa=mean(
            abs(value) for value in pressure_differences
        ),
        maximum_absolute_pressure_difference_hpa=max_pressure,
        maximum_absolute_time_difference_seconds=max_time,
        pressure_within_tolerance=max_pressure <= pressure_tolerance_hpa,
        time_within_tolerance=max_time <= time_tolerance_seconds,
        rows=tuple(rows),
    )


def suggested_pressure_offset(validation: ValidationSummary) -> float:
    """Geef de offset die de gemiddelde drukafwijking compenseert."""
    return -validation.mean_pressure_difference_hpa


def _style_header(worksheet) -> None:
    """Pas consistente kopopmaak toe."""
    fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def build_excel(
    file_data: bytes,
    input_filename: str,
    pressure_scale: float = DEFAULT_PRESSURE_SCALE,
    pressure_offset_hpa: float = DEFAULT_PRESSURE_OFFSET_HPA,
    absolute_data_offset: int | None = None,
    start_datetime_override: datetime | None = None,
    datetime_field_index: int = 0,
    reference: tuple[ReferenceMeasurement, ...] | None = None,
    pressure_tolerance_hpa: float = 0.1,
    time_tolerance_seconds: float = 1.0,
) -> ExcelResult:
    """Converteer loggerbytes naar een gevalideerde Excel-werkmap in geheugen."""
    metadata = detect_metadata(
        file_data=file_data,
        absolute_data_offset=absolute_data_offset,
        start_datetime_override=start_datetime_override,
        datetime_field_index=datetime_field_index,
    )
    measurements = tuple(
        iter_measurements(file_data, metadata, pressure_scale, pressure_offset_hpa)
    )
    if not measurements:
        raise ValueError("Geen volledige meetrecords gevonden.")
    if len(measurements) + 1 > EXCEL_MAX_ROWS:
        raise ValueError("Het aantal metingen overschrijdt de Excel-rijlimiet.")
    validation = (
        compare_measurements(
            measurements,
            reference,
            pressure_tolerance_hpa,
            time_tolerance_seconds,
        )
        if reference
        else None
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metingen"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(["Datum/tijd", "Druk (hPa)", "Ruwe druk", "Ruwe temperatuur"])
    _style_header(sheet)
    for measurement in measurements:
        sheet.append(
            [
                measurement.timestamp,
                measurement.pressure_hpa,
                measurement.raw_pressure,
                measurement.raw_temperature,
            ]
        )
        sheet.cell(sheet.max_row, 1).number_format = "dd-mm-yyyy hh:mm:ss"
        sheet.cell(sheet.max_row, 2).number_format = "0.000"
    for column, width in {"A": 22, "B": 16, "C": 14, "D": 18}.items():
        sheet.column_dimensions[column].width = width
    table = Table(displayName="TabelMetingen", ref=f"A1:D{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    sheet.add_table(table)

    metadata_sheet = workbook.create_sheet("Metadata")
    metadata_rows = [
        ("Eigenschap", "Waarde"),
        ("Invoerbestand", Path(input_filename).name),
        ("Startdatum/tijd gebruikt", metadata.start_datetime),
        ("Meetinterval", str(metadata.sample_interval)),
        ("Aantal metingen", len(measurements)),
        ("Begin meetdata, byte-offset", metadata.data_offset),
        ("Ruwe druk naar cmH2O", pressure_scale),
        ("Drukoffset (hPa)", pressure_offset_hpa),
        ("cmH2O naar hPa", CMH2O_TO_HPA),
        ("Drukberekening", "ruwe druk × schaalfactor × 0,980665 + offset"),
        ("Gevonden datumvelden", "; ".join(value.isoformat(" ") for value in metadata.detected_datetimes)),
    ]
    for row in metadata_rows:
        metadata_sheet.append(row)
    _style_header(metadata_sheet)
    metadata_sheet.freeze_panes = "A2"
    metadata_sheet.column_dimensions["A"].width = 32
    metadata_sheet.column_dimensions["B"].width = 70
    metadata_sheet["B3"].number_format = "dd-mm-yyyy hh:mm:ss"

    if validation:
        validation_sheet = workbook.create_sheet("Referentiecontrole")
        validation_sheet.append(
            [
                "Regel", "Berekende datum/tijd", "Referentie datum/tijd",
                "Tijdverschil (s)", "Berekende druk (hPa)",
                "Referentiedruk (hPa)", "Drukverschil (hPa)"
            ]
        )
        _style_header(validation_sheet)
        validation_sheet.freeze_panes = "A2"
        for row in validation.rows:
            validation_sheet.append(
                [
                    row.index, row.calculated_timestamp, row.reference_timestamp,
                    row.time_difference_seconds, row.calculated_pressure_hpa,
                    row.reference_pressure_hpa, row.pressure_difference_hpa
                ]
            )
            for column in (2, 3):
                validation_sheet.cell(validation_sheet.max_row, column).number_format = "dd-mm-yyyy hh:mm:ss"
            for column in (4, 5, 6, 7):
                validation_sheet.cell(validation_sheet.max_row, column).number_format = "0.000"
        for column, width in {"A": 10, "B": 22, "C": 22, "D": 18, "E": 22, "F": 22, "G": 22}.items():
            validation_sheet.column_dimensions[column].width = width

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    content = output.getvalue()
    check = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if check["Metingen"].max_row - 1 != len(measurements):
            raise ValueError("Validatie van het aantal Excel-meetregels is mislukt.")
    finally:
        check.close()
    return ExcelResult(
        filename=f"{Path(input_filename).stem}_hpa.xlsx",
        content=content,
        measurements=measurements,
        metadata=metadata,
        validation=validation,
    )
