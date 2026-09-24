"""Tests voor de converter."""

from datetime import datetime, timedelta
from io import BytesIO
import struct

import pytest
from openpyxl import load_workbook

from converter import build_excel, detect_metadata, raw_pressure_to_hpa


def sample_file(records: list[tuple[int, int]] | None = None) -> bytes:
    """Maak representatieve loggerbytes met twee datumvelden."""
    header = b"Logger 00:00:12 17/03/26 interval 00 02:00:00 0T end 00:00:12 17/03/26"
    payload = b"".join(struct.pack("<HH", *record) for record in (records or [(100, 200)]))
    return header + payload


def test_detect_metadata() -> None:
    metadata = detect_metadata(sample_file())
    assert metadata.start_datetime == datetime(2026, 3, 17, 0, 0, 12)
    assert metadata.sample_interval == timedelta(hours=2)


def test_pressure_conversion() -> None:
    assert raw_pressure_to_hpa(100, 0.1) == pytest.approx(9.80665)


def test_build_excel_contains_measurements_and_metadata() -> None:
    result = build_excel(sample_file([(100, 1), (200, 2)]), "logger.dat")
    workbook = load_workbook(BytesIO(result.content), data_only=True)
    try:
        assert workbook.sheetnames == ["Metingen", "Metadata"]
        assert workbook["Metingen"].max_row == 3
        assert workbook["Metingen"]["B2"].value == pytest.approx(9.80665)
        assert result.measurement_count == 2
    finally:
        workbook.close()


def test_empty_input_is_rejected() -> None:
    with pytest.raises(ValueError, match="leeg"):
        build_excel(b"", "leeg.dat")


def test_invalid_header_is_rejected() -> None:
    with pytest.raises(ValueError, match="Geen datum/tijdvelden"):
        build_excel(b"geen geldige header", "fout.dat")


def test_zero_pressure_scale_is_rejected() -> None:
    with pytest.raises(ValueError, match="groter zijn dan nul"):
        build_excel(sample_file(), "logger.dat", pressure_scale=0)
